from __future__ import annotations

import csv
import io
from typing import Iterable
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.customer import Customer
from app.models.integrations import FileAsset, ImportRun
from app.services.storage_service import get_storage

CUSTOMER_COLUMNS = {"name", "email", "phone", "tax_id", "status", "source", "notes"}


def _rows_from_bytes(filename: str, data: bytes) -> Iterable[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = data.decode("utf-8-sig")
        yield from csv.DictReader(io.StringIO(text))
        return
    if lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is required for XLSX imports") from exc
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [str(v).strip() if v is not None else "" for v in next(rows)]
        except StopIteration:
            return
        for row in rows:
            yield {headers[i]: "" if i >= len(row) or row[i] is None else str(row[i]).strip() for i in range(len(headers))}
        return
    raise ValueError("Only CSV and XLSX are supported")


def validate_customer_rows(rows: Iterable[dict[str, str]], max_rows: int = 10000) -> tuple[list[dict[str, str]], list[dict]]:
    valid: list[dict[str, str]] = []
    errors: list[dict] = []
    for number, raw in enumerate(rows, start=2):
        if number > max_rows + 1:
            errors.append({"row": number, "error": "Maximum row limit exceeded"})
            break
        normalized = {str(k).strip().lower(): (str(v).strip() if v is not None else "") for k, v in raw.items() if k is not None}
        unknown = sorted(set(normalized) - CUSTOMER_COLUMNS)
        name = normalized.get("name", "")
        email = normalized.get("email") or None
        if unknown:
            errors.append({"row": number, "error": f"Unknown columns: {', '.join(unknown)}"})
            continue
        if not name:
            errors.append({"row": number, "error": "name is required"})
            continue
        if email and "@" not in email:
            errors.append({"row": number, "error": "invalid email"})
            continue
        normalized.setdefault("status", "active")
        valid.append(normalized)
    return valid, errors


def run_customer_import(db: Session, import_run_id: UUID) -> ImportRun:
    run = db.get(ImportRun, import_run_id)
    if run is None:
        raise ValueError("Import run not found")
    storage = get_storage()
    data = storage.get(run.file_asset.storage_key)
    run.status = "running"
    run.started_at = __import__("datetime").datetime.now(__import__("datetime").timezone.utc)
    db.commit()
    rows = list(_rows_from_bytes(run.file_asset.filename, data))
    valid, errors = validate_customer_rows(rows)
    run.total_rows = len(rows)
    run.valid_rows = len(valid)
    run.error_rows = len(errors)
    run.errors = errors[:200]
    if run.mode == "dry_run":
        run.status = "completed"
        run.completed_at = __import__("datetime").datetime.now(__import__("datetime").timezone.utc)
        db.commit(); db.refresh(run)
        return run
    imported = 0
    skipped = 0
    for row in valid:
        email = row.get("email")
        existing = None
        if email:
            existing = db.scalar(select(Customer).where(Customer.company_id == run.company_id, Customer.email == email.lower()))
        if existing:
            skipped += 1
            continue
        db.add(Customer(company_id=run.company_id, name=row["name"], email=email.lower() if email else None, phone=row.get("phone") or None, tax_id=row.get("tax_id") or None, status=row.get("status") or "active", source=row.get("source") or None, notes=row.get("notes") or None))
        imported += 1
    run.imported_rows = imported
    run.skipped_rows = skipped
    run.status = "completed"
    run.completed_at = __import__("datetime").datetime.now(__import__("datetime").timezone.utc)
    db.commit(); db.refresh(run)
    return run
